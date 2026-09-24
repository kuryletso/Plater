"""Regressions found by ingesting real Google Docs / Word templates.

Every ``xfail(strict=True)`` here is an open task from the 2026-09-04 manual test
pass. Fixing the task turns the test into an XPASS *failure*, which is the signal
to drop the marker — so this module doubles as the checklist.

See tests/fixtures/real_templates/README.md for what each file exercises.
"""

import io
from collections import Counter
import zipfile

from lxml import etree
import pytest

from app.core.diagnostics import DiagnosticCollector
from app.document_engine.blueprint.models.paragraph import ParagraphBlueprint
from app.document_engine.blueprint.models.segment import (
    GroupedPlaceholderSegment,
    JoinedPlaceholderSegment,
    PlaceholderSegment,
    TextSegment,
)
from app.document_engine.blueprint.models.table import (
    CellBlueprint,
    CellPlaceholder,
    TableBlueprint,
    TablePlaceholder,
)
from app.document_engine.enums.enums import PlaceholderType
from app.document_engine.orchestration.pipeline import TemplateRenderingPipeline
from app.document_engine.parser.models.blocks import ParagraphNode
from app.document_engine.parser.parser import DocxParser
from app.document_engine.normalization.normalizers.paragraphs import normalize_paragraph
from app.document_engine.rendering.context import (
    InvoiceLineRow,
    InvoiceTableData,
    RenderContext,
)
from app.document_engine.rendering.docx.emitter import DocxEmitter
from app.document_engine.rendering.resolve.resolver import DocumentResolver


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TASK = "open task from the 2026-09-04 manual test pass"


class NoAssets:
    def get(self, asset_id):
        return None


# --- blueprint walkers -------------------------------------------------------

def _blocks(blueprint):
    """Every block in the document, tables and cells flattened in."""

    def walk(blocks):
        for block in blocks:
            yield block
            if isinstance(block, TableBlueprint):
                for row in block.rows:
                    for cell in row.cells:
                        if isinstance(cell, CellBlueprint):
                            yield from walk(cell.blocks)

    for section in blueprint.sections:
        yield from walk(section.blocks)
        for group in (section.headers, section.footers):
            for hf in (group.default, group.first, group.even):
                if hf is not None:
                    yield from walk(hf.blocks)


def segments(blueprint):
    """Top-level segments of every paragraph, in document order."""
    for block in _blocks(blueprint):
        if isinstance(block, ParagraphBlueprint):
            yield from block.segments


def placeholders(blueprint):
    """Every PlaceholderSegment, including those nested in joined/grouped forms."""
    for segment in segments(blueprint):
        if isinstance(segment, PlaceholderSegment):
            yield segment
        elif isinstance(segment, JoinedPlaceholderSegment):
            yield from (i for i in segment.items if isinstance(i, PlaceholderSegment))
        elif isinstance(segment, GroupedPlaceholderSegment):
            for group in segment.items:
                yield from (i for i in group if isinstance(i, PlaceholderSegment))


def cell_placeholders(blueprint):
    for block in _blocks(blueprint):
        if isinstance(block, TableBlueprint):
            for row in block.rows:
                yield from (c for c in row.cells if isinstance(c, CellPlaceholder))


def texts(blueprint) -> list[str]:
    return [s.text for s in segments(blueprint) if isinstance(s, TextSegment)]


def codes(diagnostics) -> list[str]:
    return [item.code for item in diagnostics.items]


# --- render helpers ----------------------------------------------------------

def _document_xml(docx: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def render_raw_xml(blueprint) -> etree._Element:
    """Raw-render (KEYS mode) and hand back word/document.xml."""
    result = TemplateRenderingPipeline(NoAssets()).render_raw(blueprint)
    assert result.docx is not None
    return _document_xml(result.docx)


def render_values_xml(blueprint, context: RenderContext) -> etree._Element:
    """Resolve + emit directly, skipping validate_context.

    The pipeline's gate refuses a context with anything missing, which would make
    these assertions depend on unrelated open tasks. Everything downstream of the
    gate is the same code path.
    """

    resolved = DocumentResolver(context, NoAssets(), DiagnosticCollector()).resolve(blueprint)
    return _document_xml(DocxEmitter(DiagnosticCollector()).emit(resolved))


def every_value_context(blueprint, *, language: str = "ENG") -> RenderContext:
    """A context shaped the way InvoiceMapper shapes one: scalars and one table.

    COLUMN placeholders stay out of the scalars, as they do in the mapper; one
    outside a table row must render empty rather than raise (Task 6).
    """

    keys = {
        p.key for p in placeholders(blueprint)
        if p.ph_type is PlaceholderType.SCALAR
    }

    columns = (
        "invl_n", "invl_desc", "invl_unit",
        "invl_qnty", "invl_price", "invl_tax", "invl_total",
    )
    row = InvoiceLineRow(values={key: {language: f"<{key}>"} for key in columns})

    return RenderContext(
        scalars={key: {language: f"<{key}>"} for key in keys},
        table=InvoiceTableData(
            rows=(row, row),
            show_tax=True,
            subtotal={language: "100.00"},
            total_tax={language: "20.00"},
            total={language: "120.00"},
            labels={key: {language: key} for key in (*columns, "subtotal", "total_tax", "total")},
        ),
    )


def grids(document) -> list[list[int]]:
    return [
        [int(float(col.get(f"{W}w", "0"))) for col in grid.findall(f"{W}gridCol")]
        for grid in document.iter(f"{W}tblGrid")
    ]


# --- baseline: what already works -------------------------------------------

def test_empty_template_ingests_without_diagnostics(ingest_real):
    blueprint, diagnostics = ingest_real("empty")

    assert not diagnostics.has_errors
    assert list(placeholders(blueprint)) == []


def test_language_suffixes_resolve_to_concrete_codes(ingest_real):
    blueprint, _ = ingest_real("languages", primary="ENG", secondary="UKR")
    dates = [p for p in placeholders(blueprint) if p.key == "date"]

    # bare {{ date }} bakes in the primary language at ingestion
    assert {p.language for p in dates} == {"ENG", "UKR"}
    assert sum(p.language == "ENG" for p in dates) == 2


def test_unknown_keys_and_languages_degrade_to_literal_text(ingest_real):
    blueprint, diagnostics = ingest_real("languages_invalid", primary="ENG", secondary="UKR")

    assert not diagnostics.has_errors, "a bad placeholder must warn, never block import"
    assert "{{ foobar }}" in texts(blueprint)
    assert "{{ date.SPA }}" in texts(blueprint)
    assert {p.key for p in placeholders(blueprint)} <= {
        "date", "prefix", "client_name", "client_country", "client_address", "client_phone",
    }


def test_placeholder_split_by_a_line_break_still_resolves(ingest_real):
    """Shift+Enter inside {{ }} works today. Paragraph breaks do not (see below)."""
    blueprint, _ = ingest_real("languages_invalid", primary="ENG", secondary="UKR")

    assert any(p.key == "prefix" for p in placeholders(blueprint))


def test_invoice_table_becomes_a_table_placeholder_at_every_nesting_depth(ingest_real):
    """Five, not six: {{ invoice_table }} in one cell of a 3x3 table replaces the
    whole wrapper table, which is what _promote_placeholder_rows is for.
    """

    blueprint, _ = ingest_real("tables")
    tables = [b for b in _blocks(blueprint) if isinstance(b, TablePlaceholder)]

    assert len(tables) == 5
    assert all(t.language == "ENG" for t in tables)


def test_invoice_line_cells_become_cell_placeholders(ingest_real):
    blueprint, _ = ingest_real("tables")

    assert {c.key for c in cell_placeholders(blueprint)} == {
        "invl_n", "invl_desc", "invl_unit", "invl_price",
    }


# --- Task 4: escapes in a placeholder separator ------------------------------

def test_escaped_newline_in_a_separator_is_a_newline(ingest_real):
    r"""`sep="\n"` must join with a line break, not the letter n.

    The same template also uses `sep=" | "` and `sep=" ::: "`, which must come
    through untouched — the escape map has to be a lookup, not a rewrite.
    """

    blueprint, _ = ingest_real("formatting")
    separators = {
        segment.separator for segment in segments(blueprint)
        if isinstance(segment, (GroupedPlaceholderSegment, JoinedPlaceholderSegment))
    }

    assert "\n" in separators
    assert "n" not in separators
    assert {" | ", " ::: "} <= separators


def test_an_escaped_newline_separator_emits_a_line_break(ingest_real):
    """End of the chain: the run builder splits on \\n and writes <w:br/>."""

    blueprint, _ = ingest_real("formatting")
    document = render_values_xml(
        blueprint, every_value_context(blueprint),
    )

    soft_breaks = [
        element for element in document.iter(f"{W}br")
        if element.get(f"{W}type") is None
    ]
    assert soft_breaks


def test_tokenizer_maps_the_standard_escapes():
    from app.document_engine.blueprint.builders.tokenizer import TK, tokenize_placeholder

    tokens = tokenize_placeholder(r'a, b, sep="\n\t\\\"x"')
    strings = [t.value for t in tokens if t.kind == TK.STRING]

    assert strings == ['\n\t\\"x']


# --- Task 5: the split-brace / multi-colour placeholder ----------------------

def test_placeholder_split_across_two_runs_at_the_braces(make_runs):
    """A colour change between `{` and `{` must not hide the placeholder.

    Real editors split runs per character when text is painted letter by letter,
    so the opening brace itself lands in a run of its own.
    """

    path = make_runs([("{", False), ("{ org_name }}", True)])
    with DocxParser(path, diagnostics=DiagnosticCollector()) as parser:
        parsed = parser.parse()

    paragraph = normalize_paragraph(
        next(block for block in parsed if isinstance(block, ParagraphNode))
    )

    assert len(paragraph.inlines) == 1
    assert paragraph.inlines[0].text == "{{ org_name }}"


def test_a_multicoloured_placeholder_resolves(ingest_real):
    """formatting.docx paints one {{ client_name }} a letter per colour."""
    blueprint, _ = ingest_real("formatting")

    assert "{" not in "".join(texts(blueprint))
    # 4 alignment groups + 1 table group + 3 merged cells + 4 font lines
    # + 4 formatting lines; one of the font lines is the rainbow one
    assert sum(p.key == "client_name" for p in placeholders(blueprint)) == 16


# --- Task 6: invoice-line placeholders outside a table -----------------------

def test_a_column_placeholder_outside_a_table_warns_at_ingestion(ingest_real):
    """{{ invl_desc }} in a plain paragraph cannot expand; say so on import."""
    _, diagnostics = ingest_real("tables")

    assert "column_placeholder_outside_table" in codes(diagnostics)
    assert not diagnostics.has_errors


def test_a_column_placeholder_outside_a_table_does_not_block_rendering(ingest_real):
    """This is what stopped TEST-004 generating: validate_context files a COLUMN
    placeholder under `scalars`, finds nothing there, and errors.
    """

    from app.document_engine.rendering.validate import validate_context

    blueprint, _ = ingest_real("tables")
    diagnostics = DiagnosticCollector()
    validate_context(blueprint, every_value_context(blueprint), diagnostics)

    missing = [
        item for item in diagnostics.items
        if item.code == "missing_required_value"
        and (item.context or {}).get("key", "").startswith("invl_")
    ]
    assert missing == []


def test_column_placeholders_keep_their_type_through_ingestion(ingest_real):
    """Guard on the fix above: they must stay COLUMN, not be rewritten to SCALAR."""
    blueprint, _ = ingest_real("tables")
    standalone = [p for p in placeholders(blueprint) if p.key.startswith("invl_")]

    assert standalone
    assert all(p.ph_type is PlaceholderType.COLUMN for p in standalone)


# --- Task 7: titlePg must mirror the source ----------------------------------

def test_titlepg_is_not_invented(ingest_real):
    """layout.docx declares a `first` header but no <w:titlePg/>, so Word shows the
    default header on page 1. Emitting titlePg blanks the first page instead.
    """

    blueprint, _ = ingest_real("layout")
    document = render_raw_xml(blueprint)

    assert document.find(f".//{W}titlePg") is None


def test_titlepg_survives_when_the_source_declares_it(tmp_path, fixture_provider):
    """The other direction: a genuine "different first page" must be kept."""
    from docx import Document
    from app.document_engine.orchestration.pipeline import TemplateIngestionPipeline

    document = Document()
    document.add_paragraph("Invoice for {{ org_name }}")
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "First page only"
    path = tmp_path / "title_page.docx"
    document.save(path)

    pipeline = TemplateIngestionPipeline(fixture_provider)
    blueprint = pipeline.finalize(pipeline.ingest(path).draft)

    assert blueprint.sections[-1].style.title_page is True
    assert render_raw_xml(blueprint).find(f".//{W}titlePg") is not None


# --- Task 9: the system invoice table's column widths ------------------------

def test_a_standalone_invoice_table_fits_the_page(ingest_real):
    blueprint, _ = ingest_real("tables")
    document = render_values_xml(
        blueprint, every_value_context(blueprint),
    )

    section = blueprint.sections[0].style
    usable = section.page_width - section.margins.left - section.margins.right

    for widths in grids(document):
        assert sum(widths) <= usable, (
            f"table grid is {sum(widths)} twips, page allows {usable}"
        )


def test_the_invoice_table_does_not_use_equal_column_widths(ingest_real):
    """A '#' column as wide as 'Description' is what equal widths look like."""
    blueprint, _ = ingest_real("tables")
    document = render_values_xml(
        blueprint, every_value_context(blueprint),
    )

    seven = [widths for widths in grids(document) if len(widths) == 7]

    assert seven, "the value render builds the full seven-column invoice table"
    assert any(len(set(widths)) > 1 for widths in seven)


# --- Task 10: page breaks ----------------------------------------------------

def test_a_page_break_survives_to_the_rendered_document(ingest_real):
    blueprint, _ = ingest_real("formatting")
    document = render_raw_xml(blueprint)

    breaks = [
        element for element in document.iter(f"{W}br")
        if element.get(f"{W}type") == "page"
    ]
    assert len(breaks) == 4         # formatting.docx has exactly four


# --- Task 11: line spacing ---------------------------------------------------

def test_line_spacing_survives(ingest_real):
    """The source sets 2.0 line spacing on a run of paragraphs; we emit single."""
    blueprint, _ = ingest_real("formatting")
    document = render_raw_xml(blueprint)

    lines = {
        element.get(f"{W}line")
        for element in document.iter(f"{W}spacing")
        if element.get(f"{W}line")
    }
    # 1.15 inherited from docDefaults (what the report was about), and the four
    # "spacing 2.00" paragraphs set directly
    assert {"276", "480"} <= lines


# --- Task 12: table alignment ------------------------------------------------

def test_table_alignment_survives(ingest_real):
    """A right-aligned table came out left-aligned; <w:jc> is never emitted."""
    blueprint, _ = ingest_real("formatting")
    document = render_raw_xml(blueprint)

    aligned = [
        element for element in document.iter(f"{W}tblPr")
        if element.find(f"{W}jc") is not None
    ]
    values = Counter(element.find(f"{W}jc").get(f"{W}val") for element in aligned)
    assert values["right"] == 2     # the half-width table and the one beside it


# --- Task 22 (post-1.0): placeholders split by a paragraph break -------------

@pytest.mark.xfail(strict=True, reason="Task 22 — deferred post-1.0 ergonomics")
def test_a_placeholder_split_by_a_paragraph_break_resolves(ingest_real):
    """Writing a grouped placeholder across Enter presses.

    Shift+Enter already works (see above), so this is convenience, not correctness.
    """

    blueprint, diagnostics = ingest_real("languages_invalid", primary="ENG", secondary="UKR")

    assert "unclosed_placeholder" not in codes(diagnostics)
    assert "{{ date" not in "".join(texts(blueprint))


# --- found while landing tasks 5-9 --------------------------------------------

def test_a_header_containing_a_table_ingests(tmp_path, fixture_provider):
    """The header/footer path called normalize_table() without its diagnostics
    argument, so any template with a table in a header or footer failed to import."""
    from docx import Document
    from docx.shared import Inches
    from app.document_engine.orchestration.pipeline import TemplateIngestionPipeline

    document = Document()
    document.add_paragraph("Body")
    table = document.sections[0].header.add_table(1, 2, Inches(6))
    table.cell(0, 0).text = "{{ org_name }}"
    table.cell(0, 1).text = "Header"
    path = tmp_path / "header_table.docx"
    document.save(path)

    result = TemplateIngestionPipeline(fixture_provider).ingest(path)

    assert not result.diagnostics.has_errors


def test_a_blueprint_stored_before_title_page_existed_still_loads(ingest_real):
    """Every blueprint already in a user's database predates `title_page` and
    `engine_version`. A new field without a default makes all of them unloadable."""
    from app.document_engine.blueprint.serialize import dump_blueprint, load_blueprint

    blueprint, _ = ingest_real("empty")
    sections, placeholder_defs, config = dump_blueprint(blueprint)
    for section in sections:
        del section["style"]["title_page"]
    del config["engine_version"]

    loaded = load_blueprint(sections, placeholder_defs, config)

    assert all(section.style.title_page is False for section in loaded.sections)
    assert loaded.config.engine_version == 0


# --- tasks 10-12: cases the fixtures do not isolate ---------------------------

def _ingest_docx(path, provider):
    from app.document_engine.orchestration.pipeline import TemplateIngestionPipeline

    pipeline = TemplateIngestionPipeline(provider)
    return pipeline.finalize(pipeline.ingest(path).draft)


def test_run_parts_split_on_hard_breaks_only():
    """A soft break stays text, since a placeholder may span one."""
    from app.document_engine.enums.enums import BreakType
    from app.document_engine.parser.extractors.runs import extract_run_parts

    run = etree.fromstring(
        f'<w:r xmlns:w="{W[1:-1]}"><w:t>a</w:t><w:br/><w:t>b</w:t>'
        f'<w:br w:type="column"/><w:t>c</w:t><w:br w:type="page"/></w:r>'
    )

    assert extract_run_parts(run) == ["a\nb", BreakType.COLUMN, "c", BreakType.PAGE]


def test_a_page_break_keeps_its_place_inside_a_run(tmp_path, fixture_provider):
    """Word writes 'Before', the break and 'After' into a single run."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from app.document_engine.enums.enums import BreakType

    document = Document()
    run = document.add_paragraph().add_run("Before {{ org_name }}")
    run.add_break(WD_BREAK.PAGE)
    run.add_text("After")
    path = tmp_path / "break.docx"
    document.save(path)

    blueprint = _ingest_docx(path, fixture_provider)
    parts = blueprint.sections[0].blocks[0].segments

    assert [type(s).__name__ for s in parts] == [
        "TextSegment", "PlaceholderSegment", "BreakSegment", "TextSegment",
    ]
    assert parts[2].kind is BreakType.PAGE

    paragraph = render_raw_xml(blueprint).find(f"{W}body/{W}p")
    order = [
        ("br", child.get(f"{W}type")) if child.tag == f"{W}br" else ("t", child.text)
        for run_ in paragraph.iter(f"{W}r")
        for child in run_
        if child.tag in (f"{W}t", f"{W}br")
    ]
    assert order == [("t", "Before "), ("t", "{{ org_name }}"), ("br", "page"), ("t", "After")]


def test_page_break_before_survives(tmp_path, fixture_provider):
    from docx import Document

    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("Second").paragraph_format.page_break_before = True
    path = tmp_path / "page_break_before.docx"
    document.save(path)

    body = render_raw_xml(_ingest_docx(path, fixture_provider)).find(f"{W}body")
    flags = [p.find(f"{W}pPr/{W}pageBreakBefore") is not None for p in body.findall(f"{W}p")]

    assert flags == [False, True]


def test_an_explicitly_false_page_break_before_is_not_emitted(ingest_real):
    """formatting.docx carries <w:pageBreakBefore w:val="0"/> on 76 paragraphs.
    Read as true, every one of them would start a new page."""

    document = render_raw_xml(ingest_real("formatting")[0])

    assert document.find(f".//{W}pageBreakBefore") is None


@pytest.mark.parametrize("val, expected", [
    (None, True), ("1", True), ("true", True), ("on", True),
    ("0", False), ("false", False), ("off", False),
])
def test_on_off_values_follow_st_onoff(val, expected):
    """Google Docs writes w:val="0"; LibreOffice writes "false"."""
    from app.document_engine.parser.extractors.styles import has_tag

    attr = f' w:val="{val}"' if val is not None else ""
    ppr = etree.fromstring(f'<w:pPr xmlns:w="{W[1:-1]}"><w:pageBreakBefore{attr}/></w:pPr>')

    assert has_tag(ppr, "w:pageBreakBefore") is expected


@pytest.mark.parametrize("val, expected", [("start", "left"), ("end", "right"), ("center", "center")])
def test_start_and_end_alignments_read_as_left_and_right(val, expected):
    """Unmapped, start/end reach the enum and raise a bare ValueError mid-import."""
    from app.document_engine.parser.extractors.styles import (
        extract_paragraph_style, extract_table_style,
    )

    namespace = f'xmlns:w="{W[1:-1]}"'
    ppr = etree.fromstring(f'<w:pPr {namespace}><w:jc w:val="{val}"/></w:pPr>')
    tbl_pr = etree.fromstring(f'<w:tblPr {namespace}><w:jc w:val="{val}"/></w:tblPr>')

    assert extract_paragraph_style(ppr).alignment == expected
    assert extract_table_style(tbl_pr).alignment == expected


def test_line_spacing_is_inherited_from_the_normal_style(tmp_path, fixture_provider):
    """Word keeps body spacing on the Normal style. A paragraph without a pStyle
    uses it — exactly as runs already inherit its font size."""
    from docx import Document

    document = Document()
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.add_paragraph("Body")
    path = tmp_path / "normal_spacing.docx"
    document.save(path)

    body = render_raw_xml(_ingest_docx(path, fixture_provider)).find(f"{W}body")
    spacing = body.find(f"{W}p/{W}pPr/{W}spacing")

    assert (spacing.get(f"{W}line"), spacing.get(f"{W}lineRule")) == ("276", "auto")


def test_table_jc_sits_where_the_schema_puts_it(ingest_real):
    """CT_TblPr is ordered: jc directly after tblW. Out of order, Word may refuse the file."""
    document = render_raw_xml(ingest_real("formatting")[0])

    for tbl_pr in document.iter(f"{W}tblPr"):
        names = [etree.QName(child).localname for child in tbl_pr]
        assert names.index("jc") == names.index("tblW") + 1


# --- tabs --------------------------------------------------------------------

def run_contents(body) -> list[str]:
    """Run children in order: text as its own string, anything else by tag name."""

    return [
        child.text if child.tag == f"{W}t" else etree.QName(child).localname
        for run in body.iter(f"{W}r")
        for child in run
        if child.tag != f"{W}rPr"
    ]


def test_a_tab_is_emitted_as_an_element(tmp_path, fixture_provider):
    """A tab character inside <w:t> is only whitespace to Word, which drops it.
    A tab has to be <w:tab/> — text falling back from a placeholder showed this."""

    from docx import Document

    document = Document()
    run = document.add_paragraph().add_run("before")
    run.add_tab()
    run.add_text("after")
    path = tmp_path / "tabbed.docx"
    document.save(path)

    body = render_raw_xml(_ingest_docx(path, fixture_provider)).find(f"{W}body")

    assert run_contents(body) == ["before", "tab", "after"]


def test_tabs_survive_in_text_a_placeholder_fell_back_to(tmp_path, fixture_provider):
    """An unclosed placeholder renders literally — tabs included."""

    from docx import Document

    document = Document()
    run = document.add_paragraph().add_run("{{ org_name")
    run.add_tab()
    run.add_text("unclosed")
    path = tmp_path / "fallback.docx"
    document.save(path)

    body = render_raw_xml(_ingest_docx(path, fixture_provider)).find(f"{W}body")

    assert run_contents(body) == ["{{ org_name", "tab", "unclosed"]


def test_a_tab_and_a_line_break_in_one_run_keep_their_order(tmp_path, fixture_provider):
    """Both are elements, so the split has to interleave them, not do one then the other."""

    from docx import Document

    document = Document()
    run = document.add_paragraph().add_run("a")
    run.add_tab()
    run.add_text("b")
    run.add_break()
    run.add_tab()
    run.add_text("c")
    path = tmp_path / "tab_and_break.docx"
    document.save(path)

    body = render_raw_xml(_ingest_docx(path, fixture_provider)).find(f"{W}body")

    assert run_contents(body) == ["a", "tab", "b", "br", "tab", "c"]


# --- Task 16: images in headers and footers ----------------------------------

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


class BundleAssets:
    """An AssetProvider over an ingestion bundle, so rendering keeps the images."""

    def __init__(self, bundle):
        self._bundle = bundle

    def get(self, asset_id):
        from app.document_engine.rendering.ports import Asset

        blob = self._bundle.get(asset_id)
        return Asset(data=blob.data, mime=blob.mime_type) if blob is not None else None


def real_provider():
    from app.document_engine.blueprint.models.template import TemplateConfig
    from tests.conftest import FixtureInputProvider, real_placeholder_defaults

    return FixtureInputProvider(
        placeholders=real_placeholder_defaults(),
        config=TemplateConfig(
            primary_language="ENG", secondary_language=None, type="invoice",
            name="template", description="", append_currency=True,
        ),
    )


def ingest_with_assets(path, provider):
    from app.document_engine.orchestration.pipeline import TemplateIngestionPipeline

    pipeline = TemplateIngestionPipeline(provider)
    result = pipeline.ingest(path)
    return pipeline.finalize(result.draft), result


def _images(blocks):
    from app.document_engine.blueprint.models.segment import ImageSegment

    for block in blocks:
        if isinstance(block, ParagraphBlueprint):
            yield from (s for s in block.segments if isinstance(s, ImageSegment))
        elif isinstance(block, TableBlueprint):
            for row in block.rows:
                for cell in row.cells:
                    if isinstance(cell, CellBlueprint):
                        yield from _images(cell.blocks)


def header_footer_images(blueprint):
    return [
        image
        for section in blueprint.sections
        for group in (section.headers, section.footers)
        for hf in (group.default, group.first, group.even)
        if hf is not None
        for image in _images(hf.blocks)
    ]


def body_images(blueprint):
    return [image for section in blueprint.sections for image in _images(section.blocks)]


@pytest.mark.xfail(strict=True, reason=f"Task 16 — {TASK}: header/footer rIds resolve through document.xml.rels")
def test_a_footer_image_resolves_through_the_footers_own_relationships(real_template):
    """layout.docx's footer shows the body's picture, as rId1 in footer1.xml.rels.
    Through document.xml.rels, rId1 is the theme, which got stored as an "image"."""

    blueprint, result = ingest_with_assets(real_template("layout"), real_provider())
    (footer_image,) = header_footer_images(blueprint)

    # both reference word/media/image1.jpg, so they must be the same asset
    assert footer_image.asset_id in {image.asset_id for image in body_images(blueprint)}
    assert all(blob.mime_type.startswith("image/") for blob in result.assets.values())


@pytest.mark.xfail(strict=True, reason=f"Task 16 — {TASK}")
def test_a_footer_image_reaches_the_rendered_footer(real_template):
    """End of the chain: the footer part holds a picture, resolved through its own
    relationships, and the bytes behind it are a JPEG."""

    import re

    blueprint, result = ingest_with_assets(real_template("layout"), real_provider())
    docx = TemplateRenderingPipeline(BundleAssets(result.assets)).render_raw(blueprint).docx

    pictures = []
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r"word/footer\d+\.xml", name):
                continue
            rids = [blip.get(R_EMBED) for blip in etree.fromstring(archive.read(name)).iter(f"{A}blip")]
            if not rids:
                continue
            rels = etree.fromstring(archive.read(f"word/_rels/{name.rsplit('/', 1)[-1]}.rels"))
            targets = {relationship.get("Id"): relationship.get("Target") for relationship in rels}
            pictures += [archive.read(f"word/{targets[rid]}")[:2] for rid in rids]

    assert pictures == [b"\xff\xd8"]        # one JPEG


@pytest.mark.xfail(strict=True, reason=f"Task 16 — {TASK}")
def test_a_header_gets_its_own_picture_not_the_bodys(tmp_path, fixture_provider):
    """Relationship ids are per part, so they collide across parts. Whatever the
    ids, a header's picture must be the header's picture."""

    from docx import Document
    from PIL import Image

    body_png, header_png = tmp_path / "body.png", tmp_path / "header.png"
    Image.new("RGB", (8, 8), "red").save(body_png)
    Image.new("RGB", (8, 8), "blue").save(header_png)

    document = Document()
    document.add_picture(str(body_png))
    document.sections[0].header.paragraphs[0].add_run().add_picture(str(header_png))
    path = tmp_path / "pictures.docx"
    document.save(path)

    blueprint, result = ingest_with_assets(path, fixture_provider)
    (header_image,) = header_footer_images(blueprint)

    assert result.assets[header_image.asset_id].data == header_png.read_bytes()


@pytest.mark.xfail(strict=True, reason=f"Task 16 — {TASK}: any related part is read as an image")
def test_a_relationship_that_is_not_an_image_is_never_read_as_one(tmp_path, fixture_provider):
    """Defence in depth: a picture pointed at the styles part is refused with a
    warning at import, not stored as an "image" that only fails at render."""

    from docx import Document
    from PIL import Image

    png = tmp_path / "body.png"
    Image.new("RGB", (8, 8), "red").save(png)
    document = Document()
    document.add_picture(str(png))
    source = tmp_path / "source.docx"
    document.save(source)

    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    by_type = {relationship.get("Type").rsplit("/", 1)[-1]: relationship.get("Id") for relationship in rels}
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        f'r:embed="{by_type["image"]}"'.encode(), f'r:embed="{by_type["styles"]}"'.encode(),
    )
    path = tmp_path / "misdirected.docx"
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    blueprint, result = ingest_with_assets(path, fixture_provider)

    assert body_images(blueprint) == []
    assert "image_relationship_not_an_image" in codes(result.diagnostics)
    assert all(blob.mime_type.startswith("image/") for blob in result.assets.values())
