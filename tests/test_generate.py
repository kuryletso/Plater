"""Generate service: assemble -> render -> consume-on-success, and the peek/preview
guarantee that only a successful generate ever advances the counter."""

from datetime import date
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy.orm import Session

from app.db.models.registries.document_type import DocumentTypeRegistry
from app.db.models.registries.placeholder import PlaceholderRegistry
from app.document_engine.enums.enums import PlaceholderType
from app.services.doc_sequence.repository import SequenceRepository
from app.services.errors import EntityNotFound, InvalidSelection, SequenceConflict
from app.services.invoice.draft import InvoiceDraft, PartySelection
from app.services.invoice.generate import InvoiceGenerateService
from app.services.template.db_input_provider import DbTemplateInputProvider
from app.services.template.import_service import TemplateImportService


@pytest.fixture
def registered_keys(session: Session) -> None:
    """Registry rows for the mapper-produced keys the test templates use."""

    session.add_all([
        PlaceholderRegistry(key=key, system=True, required=False,
                            type=PlaceholderType.SCALAR, active=True, columns=None)
        for key in ("prefix", "id", "provider_name")
    ])
    session.commit()


@pytest.fixture
def template_id(session, seeded_inputs, registered_keys, make_docx) -> int:
    """A rendering-ready template whose placeholders the mapper can all satisfy."""

    path = make_docx(
        paragraphs=["Invoice {{ prefix }}{{ id }}", "From {{ provider_name }}"],
        name="generate.docx",
    )
    service = TemplateImportService(session, DbTemplateInputProvider(session))
    return service.commit(service.ingest(path))


@pytest.fixture
def failing_template_id(session, seeded_inputs, registered_keys, make_docx) -> int:
    """A template demanding org_name, which the mapper never produces —
    ingests fine, fails at render validation."""

    path = make_docx(paragraphs=["For {{ org_name }}"], name="failing.docx")
    service = TemplateImportService(session, DbTemplateInputProvider(session))
    return service.commit(service.ingest(path))


@pytest.fixture
def scenario(session, make_org, make_line_input, make_sequence):
    """(provider, client, sequence, line) — one of everything a draft needs."""

    provider = make_org("Provider Co", tax_value="11111111")
    client = make_org("Client Co", tax_value="22222222")
    return provider, client, make_sequence(provider), make_line_input()


@pytest.fixture
def service(session: Session) -> InvoiceGenerateService:
    return InvoiceGenerateService(session)


def make_draft(template_id, scenario, **overrides) -> InvoiceDraft:
    provider, client, sequence, line = scenario
    kwargs = dict(
        template_id=template_id,
        sequence_id=sequence.id,
        currency_code="UAH",
        issue_date=date(2026, 8, 10),
        provider=PartySelection(
            organization_id=provider.id,
            tax_id_id=provider.tax_ids[0].id,
        ),
        client=PartySelection(
            organization_id=client.id,
            tax_id_id=client.tax_ids[0].id,
        ),
        lines=(line,),
    )
    kwargs.update(overrides)
    return InvoiceDraft(**kwargs)


def document_text(docx: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(docx)).paragraphs)


# --- generate ----------------------------------------------------------------

def test_generate_produces_a_document_carrying_the_issued_number(service, template_id,
                                                                 scenario):
    result = service.generate(make_draft(template_id, scenario))

    assert result.succeeded
    assert result.number.formatted == "INV-0007"
    assert "INV-0007" in document_text(result.docx)


def test_generate_consumes_exactly_one_number(service, template_id, scenario,
                                              session: Session):
    _, _, sequence, _ = scenario

    service.generate(make_draft(template_id, scenario))

    session.refresh(sequence)
    assert sequence.counter == 7


def test_consecutive_generates_number_consecutively(service, template_id, scenario):
    first = service.generate(make_draft(template_id, scenario))
    second = service.generate(make_draft(template_id, scenario))

    assert first.number.formatted == "INV-0007"
    assert second.number.formatted == "INV-0008"


def test_a_failed_render_reports_diagnostics(service, failing_template_id, scenario):
    result = service.generate(make_draft(failing_template_id, scenario))

    assert not result.succeeded
    assert result.docx is None
    assert any(item.code == "missing_required_value"
               for item in result.diagnostics.items)


def test_a_failed_render_burns_no_number(service, failing_template_id, scenario,
                                         session: Session):
    """The reason consume waits for success: a failed generate must leave the
    counter exactly where it was."""
    _, _, sequence, _ = scenario

    result = service.generate(make_draft(failing_template_id, scenario))

    assert not result.succeeded
    session.refresh(sequence)
    assert sequence.counter == 6


def test_the_number_recovers_after_a_failure(service, template_id, failing_template_id,
                                             scenario):
    """A failed attempt and the retry that follows get the same number."""
    service.generate(make_draft(failing_template_id, scenario))

    result = service.generate(make_draft(template_id, scenario))

    assert result.number.formatted == "INV-0007"


# --- preview -----------------------------------------------------------------

def test_preview_renders_without_consuming(service, template_id, scenario,
                                           session: Session):
    _, _, sequence, _ = scenario

    result = service.preview(make_draft(template_id, scenario))

    assert result.succeeded
    assert "INV-0007" in document_text(result.docx)
    session.refresh(sequence)
    assert sequence.counter == 6


def test_preview_shows_the_number_generate_will_issue(service, template_id, scenario):
    previewed = service.preview(make_draft(template_id, scenario))
    generated = service.generate(make_draft(template_id, scenario))

    assert previewed.number == generated.number


# --- preconditions -----------------------------------------------------------

def test_generate_rejects_a_sequence_of_another_document_type(service, template_id,
                                                              scenario,
                                                              session: Session):
    """Nothing else ties draft.template_id to draft.sequence_id, so an invoice
    numbered by an 'akt' sequence must be refused here."""
    provider, _, _, _ = scenario
    session.add(DocumentTypeRegistry(code="akt", system=True, active=True))
    session.commit()
    akt_sequence = SequenceRepository(session).create(provider.id, "akt", prefix="AKT-")

    with pytest.raises(InvalidSelection):
        service.generate(make_draft(template_id, scenario,
                                    sequence_id=akt_sequence.id))


def test_the_check_follows_the_template_column_not_the_ingested_config(
        service, template_id, scenario, session: Session):
    """Retyping a template must change what generation accepts. While the check
    read blueprint.config.type, an edited type was silently ignored here."""
    from app.services.template.repository import TemplateRepository

    session.add(DocumentTypeRegistry(code="akt", system=True, active=True))
    session.commit()
    TemplateRepository(session).update_metadata(template_id, document_type="akt")

    # the scenario's sequence still numbers invoices
    with pytest.raises(InvalidSelection):
        service.generate(make_draft(template_id, scenario))


def test_a_retyped_template_generates_against_a_matching_sequence(
        service, template_id, scenario, session: Session):
    from app.services.doc_sequence.repository import SequenceRepository
    from app.services.template.repository import TemplateRepository

    provider, _, _, _ = scenario
    session.add(DocumentTypeRegistry(code="akt", system=True, active=True))
    session.commit()
    TemplateRepository(session).update_metadata(template_id, document_type="akt")
    akt_sequence = SequenceRepository(session).create(provider.id, "akt", prefix="AKT-")

    result = service.generate(
        make_draft(template_id, scenario, sequence_id=akt_sequence.id)
    )

    assert result.succeeded


def test_generate_raises_for_an_unknown_sequence(service, template_id, scenario):
    with pytest.raises(EntityNotFound):
        service.generate(make_draft(template_id, scenario, sequence_id=9999))


def test_generate_raises_for_an_unknown_template(service, template_id, scenario):
    with pytest.raises(EntityNotFound):
        service.generate(make_draft(9999, scenario))


def test_a_lost_race_for_the_number_raises_instead_of_mislabeling(service, template_id,
                                                                  scenario,
                                                                  monkeypatch):
    """If another consumer takes the number between peek and consume, the
    rendered document is stale — surfacing SequenceConflict beats saving a
    document whose number was just issued to someone else."""
    real_consume = SequenceRepository.consume

    def racing_consume(self, sequence_id):
        real_consume(self, sequence_id)             # the competitor gets there first
        return real_consume(self, sequence_id)

    monkeypatch.setattr(SequenceRepository, "consume", racing_consume)

    with pytest.raises(SequenceConflict):
        service.generate(make_draft(template_id, scenario))
